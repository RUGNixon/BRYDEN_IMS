from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"C:\Users\educa\Desktop\Bryden IMS\bryden\Gisele_Presentation_Speaker_Guide.docx"


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F1F1F"
MUTED = "666666"
HEADER_FILL = "E8EEF5"
SOFT_FILL = "F4F6F9"
GOLD_FILL = "F9F4E6"


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="BFC7D1", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for i, width in enumerate(widths):
            cell = row.cells[i]
            cell.width = width
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width.inches * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, size=11, color=INK, bold=False, space_after=6, line_spacing=1.25):
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color, bold=bold)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    if level == 1:
        paragraph.style = doc.styles["Heading 1"]
    elif level == 2:
        paragraph.style = doc.styles["Heading 2"]
    else:
        paragraph.style = doc.styles["Heading 3"]
    run = paragraph.add_run(text)
    return paragraph


def add_body(doc, text, bold_label=None):
    paragraph = doc.add_paragraph()
    if bold_label:
        r = paragraph.add_run(bold_label)
        set_run_font(r, size=11, bold=True, color=INK)
        r = paragraph.add_run(text)
    else:
        r = paragraph.add_run(text)
    style_paragraph(paragraph)
    return paragraph


def add_callout(doc, title, body, fill=SOFT_FILL):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [Inches(6.5)])
    set_table_borders(table, color="D7DEE8", size="4")
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    set_cell_margins(cell, top=130, bottom=130, start=170, end=170)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    r2 = p2.add_run(body)
    set_run_font(r2, size=11, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_label_detail_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    set_table_width(table, [Inches(1.65), Inches(4.85)])
    set_table_borders(table)
    for i, (label, detail) in enumerate(rows):
        if i > 0:
            table.add_row()
        cells = table.rows[i].cells
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_fill(cells[0], HEADER_FILL)
        for cell in cells:
            set_cell_margins(cell)
        lp = cells[0].paragraphs[0]
        lp.paragraph_format.space_after = Pt(0)
        lr = lp.add_run(label)
        set_run_font(lr, size=10.5, bold=True, color=DARK_BLUE)
        dp = cells[1].paragraphs[0]
        dp.paragraph_format.space_after = Pt(0)
        dp.paragraph_format.line_spacing = 1.15
        dr = dp.add_run(detail)
        set_run_font(dr, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_guide_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    widths = [Inches(0.72), Inches(1.78), Inches(2.15), Inches(1.85)]
    set_table_width(table, widths)
    set_table_borders(table)
    headers = ["Slides", "Section", "Audience meaning", "Presenter focus"]
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_fill(cell, HEADER_FILL)
        set_cell_margins(cell, top=90, bottom=90, start=90, end=90)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.5, bold=True, color=DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_margins(cells[i], top=85, bottom=85, start=90, end=90)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(text)
            set_run_font(r, size=9.2, bold=(i == 0), color=INK if i else DARK_BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Title", 20, DARK_BLUE, 0, 8),
        ("Subtitle", 11, MUTED, 0, 12),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name.startswith("Heading") or name == "Title"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25


def add_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Presenter guide | Rwanda Dairy Development Project")
    set_run_font(r, size=8.5, color=MUTED)


def build():
    doc = Document()
    configure_styles(doc)
    add_footer(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title.add_run("Presenter Guide and Audience Summary")
    set_run_font(r, size=20, bold=True, color=DARK_BLUE)

    subtitle = doc.add_paragraph(style="Subtitle")
    r = subtitle.add_run(
        "Effect of Project Planning Practices on Performance of Projects: "
        "A Case of Rwanda Dairy Development Project in Ruhango District, Rwanda"
    )
    set_run_font(r, size=11, color=MUTED)

    add_callout(
        doc,
        "Core message",
        "Project planning is not only an administrative activity. In this study, resources planning, operational planning, and monitoring and evaluation planning all had positive and significant effects on the performance of the Rwanda Dairy Development Project in Ruhango District.",
        fill=GOLD_FILL,
    )

    add_heading(doc, "One-minute summary", 1)
    add_body(
        doc,
        "This presentation explains how planning practices affect the performance of the Rwanda Dairy Development Project in Ruhango District. The study focused on three planning practices: resources planning, operational planning, and monitoring and evaluation planning. The problem is important because public projects are expected to achieve timely completion, cost efficiency, quality standards, and proper use of resources, yet many public investments face delays, stalled work, idle assets, and quality concerns."
    )
    add_body(
        doc,
        "The research used descriptive and correlational designs with quantitative and qualitative approaches. From a target population of 365 people involved in the project, 191 respondents were sampled, and 187 completed questionnaires were returned, giving a 97.9 percent response rate. Data were analyzed using SPSS version 27, descriptive statistics, Pearson correlation, and multiple regression at the 0.05 significance level."
    )
    add_body(
        doc,
        "The major finding is that all three planning practices significantly improved project performance. Resources planning was the strongest predictor, followed by operational planning and monitoring and evaluation planning. The regression model explained 78.6 percent of the variation in project performance, meaning planning practices strongly account for how well the project performed."
    )

    add_heading(doc, "How to explain the study to the audience", 1)
    add_label_detail_table(
        doc,
        [
            ("Problem", "Public projects often struggle with delays, stalled implementation, idle assets, and quality gaps. The dairy sector also faces informal milk channels and milk quality concerns."),
            ("Gap", "Previous studies discussed project challenges, but they did not clearly test how resources planning, operational planning, and M&E planning affect this specific dairy project in Ruhango District."),
            ("Approach", "The study collected evidence from project stakeholders, tested reliability and validity, then used correlation and regression to measure the effect of planning practices on performance."),
            ("Answer", "Planning practices matter. Better resources planning, operational planning, and M&E planning are linked with better project performance."),
            ("Main action", "Project leaders should strengthen resource tracking, field coordination, beneficiary participation in M&E, and rapid hygiene action based on laboratory feedback."),
        ],
    )

    add_heading(doc, "Opening script", 1)
    add_body(
        doc,
        "Good morning. My presentation is about the effect of project planning practices on the performance of the Rwanda Dairy Development Project in Ruhango District. The study focuses on three planning practices: resources planning, operational planning, and monitoring and evaluation planning. The reason this topic matters is that public projects are expected to deliver results on time, within budget, and at the required quality. However, evidence from public investment projects and the dairy sector shows that delays, stalled projects, idle assets, informal milk marketing, and quality concerns still exist. Therefore, my study examined whether stronger planning practices can improve project performance."
    )

    add_heading(doc, "Slide-by-slide guide", 1)
    add_guide_table(
        doc,
        [
            ("1", "Title", "The audience should immediately understand the topic, case area, presenter, and date.", "Speak clearly and confidently. Do not explain too much yet."),
            ("2", "Introduction", "The study is about planning practices and project performance in Rwanda Dairy Development Project.", "Define the three planning practices in simple language."),
            ("3", "Problem", "The research matters because project delays, stalled assets, and dairy quality issues affect performance.", "Use this slide to create urgency. Mention the statistics briefly."),
            ("4", "Objectives and hypotheses", "The study tested whether each planning practice significantly affects performance.", "Explain that the hypotheses gave the study a measurable direction."),
            ("5", "Theory", "The theories explain why resources, operations, and M&E should influence results.", "Link each theory to one variable. Keep this short."),
            ("6", "Conceptual framework", "Planning practices are the independent variables, and project performance is the dependent variable.", "Show the logic: planning practices lead to performance outcomes."),
            ("7-8", "Method and quality checks", "The evidence was collected systematically and tested for validity and reliability.", "Mention sample size, tools, SPSS, CVI, and Cronbach's Alpha values above 0.70."),
            ("9-10", "Results overview", "The response rate was high, and the audience needs to know how the descriptive results are interpreted.", "Use these as preparation slides before the main findings."),
            ("11-15", "Descriptive findings", "Respondents highly agreed that planning practices and project performance were strong.", "Do not read every table. Give the mean, meaning, and one practical example."),
            ("16", "Inferential statistics", "Correlation and regression show whether relationships and effects are statistically significant.", "Explain p < 0.05 and positive relationships in plain language."),
            ("17", "Correlation", "Resources, operational, and M&E planning all had strong positive significant relationships with performance.", "Stress that all p-values were .000, below .05."),
            ("18-19", "Model strength", "The full model was strong and statistically significant.", "Highlight R Square = .786 and F = 222.629."),
            ("20", "Coefficients", "Resources planning had the strongest effect, followed by operational planning and M&E planning.", "This is the most important results slide. Slow down here."),
            ("21", "Hypotheses", "All null hypotheses were rejected because all three effects were significant.", "Say this confirms planning practices as performance drivers."),
            ("22", "Conclusion and recommendations", "The study leads to practical improvements for project leadership and stakeholders.", "Connect each recommendation to the results."),
            ("23", "Thank you", "The presentation is complete and ready for questions.", "Close confidently and invite questions."),
        ],
    )

    add_heading(doc, "Key numbers to remember", 1)
    add_label_detail_table(
        doc,
        [
            ("Response rate", "187 out of 191 questionnaires were returned complete, giving a 97.9 percent response rate."),
            ("Validity", "CVI = 0.83, above the 0.70 threshold."),
            ("Reliability", "Cronbach's Alpha values ranged from .822 to .837, all above 0.70."),
            ("Correlation", "Resources planning r = .815; M&E planning r = .775; operational planning r = .752. All were positive and significant."),
            ("Model strength", "R = .887 and R Square = .786, meaning the model explained 78.6 percent of performance variation."),
            ("Strongest predictor", "Resources planning had the largest coefficient: B = .521, t = 10.633, Sig. = .000."),
            ("Other effects", "Operational planning B = .223 and M&E planning B = .196; both were positive and significant."),
        ],
    )

    add_heading(doc, "Strong closing script", 1)
    add_body(
        doc,
        "In conclusion, this study confirms that project performance in Rwanda Dairy Development Project is strongly influenced by planning practices. Resources planning improves input readiness and allocation. Operational planning improves coordination and role clarity. Monitoring and evaluation planning improves tracking, learning, and timely correction. Since all three practices had positive and significant effects, project leaders and stakeholders should strengthen planning systems if they want better project performance in Ruhango District."
    )

    add_heading(doc, "Likely questions and strong answers", 1)
    add_label_detail_table(
        doc,
        [
            ("Why this topic?", "Because public project performance depends on planning, and evidence shows that delays, stalled projects, idle assets, and dairy quality issues remain serious concerns."),
            ("Why Ruhango District?", "Ruhango District was selected because the Rwanda Dairy Development Project operates there and provides a clear case for studying planning practices and performance."),
            ("Why these variables?", "Resources, operations, and M&E represent the main planning areas that influence whether project activities are supplied, coordinated, monitored, and corrected."),
            ("What is your strongest result?", "Resources planning had the strongest effect on performance, with B = .521 and Sig. = .000."),
            ("What is the practical implication?", "Project leaders should improve materials and vehicle tracking, strengthen weekly field coordination, involve beneficiaries in M&E, and respond quickly to laboratory feedback."),
        ],
    )

    add_heading(doc, "Presentation reminders", 1)
    add_body(doc, "Do not read every bullet. Explain the meaning first, then mention the number.")
    add_body(doc, "Move quickly through theory and methodology, then slow down on slides 17 to 21.")
    add_body(doc, "When presenting tables, say: the table shows, this means, and therefore.")
    add_body(doc, "Keep returning to the main message: better planning leads to better project performance.")

    doc.save(OUT)


if __name__ == "__main__":
    build()
